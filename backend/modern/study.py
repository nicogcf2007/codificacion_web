"""Durable surveys and frozen job plans; clustering is computed once."""
from __future__ import annotations
import csv
import io
from pathlib import PurePath
import re
from zipfile import ZipFile, BadZipFile
from sqlalchemy import select, insert, update, func
from core.clustering import cluster_survey_responses, is_empty_or_symbol
from .repository import Repository, identity, now, record_activity
from . import db


MAX_BYTES = 50 * 1024 * 1024
MAX_ROWS = 50000
csv.field_size_limit(MAX_BYTES)


def read_survey(filename, content):
    extension = PurePath(filename.replace('\\', '/')).suffix.lower()
    if not content or len(content) > MAX_BYTES:
        raise ValueError('Archivo vacío o superior a 50 MB')
    if extension == '.csv':
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = content.decode('cp1252')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        try:
            reader = csv.reader(io.StringIO(text), dialect, strict=True)
            columns = next(reader, [])
            rows = []
            for values in reader:
                if not values:
                    continue
                if len(rows) >= MAX_ROWS:
                    raise ValueError('Límite de 50.000 registros por encuesta')
                if len(values) != len(columns):
                    raise ValueError('El número de campos no coincide con el encabezado CSV')
                rows.append(dict(zip(columns, values)))
        except csv.Error as exc:
            raise ValueError('CSV inválido: revise delimitadores, comillas y límite de 50 MB por campo') from exc
    elif extension == '.xls':
        from .excel_xls import legacy_sheets
        _, values = legacy_sheets(content, max_rows=MAX_ROWS)[0]
        columns = [str(value) if value != '' else '' for value in (values[0] if values else [])]
        rows = [dict(zip(columns, row)) for row in values[1:]]
    elif extension == '.xlsx':
        from openpyxl import load_workbook
        workbook = None
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(entry.file_size for entry in archive.infolist()) > MAX_BYTES * 4:
                    raise ValueError('Excel descomprimido demasiado grande')
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            values = workbook.active.iter_rows(values_only=True)
            columns = [str(value) if value is not None else '' for value in next(values, [])]
            rows = []
            for values_row in values:
                if len(rows) >= MAX_ROWS:
                    raise ValueError('Límite de 50.000 registros por encuesta')
                row = {}
                for key, value in zip(columns, values_row):
                    if value is None:
                        row[key] = ''
                    elif hasattr(value, 'isoformat'):
                        row[key] = value.isoformat()
                    else:
                        row[key] = value
                rows.append(row)
        except (BadZipFile, OSError, KeyError, SyntaxError, ValueError, TypeError) as exc:
            raise ValueError('Excel inválido') from exc
        finally:
            if workbook is not None:
                workbook.close()
    else:
        raise ValueError('Solo se aceptan archivos .xlsx, .xls y .csv')
    if not columns or any(not name.strip() for name in columns) or len(set(columns)) != len(columns):
        raise ValueError('Encabezados vacíos o duplicados')
    if len(columns) > 500 or not rows:
        raise ValueError('La encuesta requiere registros y máximo 500 columnas')
    return columns, rows


def survey_public(item):
    return {key: value for key, value in item.items() if key not in {'original', 'rows'}} | {'preview': item['rows'][:5]}


def initial_metrics(total, prototype_count):
    return dict(api_calls=0, saved_calls=total - prototype_count,
                savings_pct=(total - prototype_count) / total * 100 if total else 0,
                cost_usd=None, saved_usd=None, purity=None, records_per_second=0,
                eta_seconds=None, input_tokens=0, output_tokens=0)



def find_code_column(column_name: str, survey_columns: list[str]) -> str | None:
    norm_map = {str(c).strip().upper(): c for c in survey_columns}
    cleaned = str(column_name).strip().upper()
    candidates = [
        f"C{cleaned}",
        f"COD_{cleaned}",
        f"{cleaned}_COD",
        f"CODIGO_{cleaned}",
        f"{cleaned}_CODIGO",
    ]
    if cleaned.startswith("P"):
        candidates.insert(1, f"CP{cleaned[1:]}")
    for cand in candidates:
        if cand in norm_map and norm_map[cand] != column_name:
            return norm_map[cand]
    return None


def extract_precoded_codes(raw_val, categories: list[dict], max_labels: int) -> list[str] | None:
    if raw_val is None:
        return None
    if isinstance(raw_val, float):
        if raw_val.is_integer():
            raw_val = int(raw_val)
        else:
            return None
    val_str = str(raw_val).strip()
    if not val_str or val_str.lower() in {'nan', 'null', 'none', '.'}:
        return None
    raw_parts = [p.strip() for p in re.split(r'[;,]', val_str) if p.strip()]
    if not raw_parts:
        return None
    allowed = {c['code'] for c in categories} | {'77', '88', '99'}
    parsed_codes = []
    for part in raw_parts:
        cand = part
        if cand in allowed:
            parsed_codes.append(cand)
        elif cand.isdigit() and cand.zfill(2) in allowed:
            parsed_codes.append(cand.zfill(2))
        else:
            return None
    if len(parsed_codes) > max_labels:
        return None
    deduped = list(dict.fromkeys(parsed_codes))
    if len(deduped) > 1 and (set(deduped) & {'88', '99'}):
        return None
    return deduped if deduped else None


class StudyRepository(Repository):
    def preview_books(self, project_id, filename, content):
        from .codebook_excel import read_codebook, public_preview
        self.require(db.projects, project_id)
        return public_preview(read_codebook(filename, content))

    def list_books(self, project_id):
        self.require(db.projects, project_id)
        with self.engine.connect() as conn:
            books = [dict(row) for row in conn.execute(select(db.codebooks).where(db.codebooks.c.project_id == project_id).order_by(db.codebooks.c.created_at.desc())).mappings()]
            sources = {row['codebook_id']: {key: value for key, value in row.items() if key != 'codebook_id'}
                       for row in conn.execute(select(db.codebook_sources).join(db.codebooks)
                                               .where(db.codebooks.c.project_id == project_id)).mappings()}
            return [{**book, 'source': sources.get(book['id'])} for book in books]

    def import_books(self, project_id, survey_id, filename, content, selections, actor_id=None):
        from .codebook_excel import read_codebook
        self.require(db.projects, project_id)
        parsed = read_codebook(filename, content)
        catalogs = {catalog['key']: catalog for catalog in parsed['catalogs']}
        result = dict(filename=parsed['filename'], codebooks=[], mappings=[])
        with self.engine.begin() as conn:
            survey = self.require(db.surveys, survey_id, conn)
            if survey['project_id'] != project_id:
                raise ValueError('La encuesta no pertenece al proyecto')
            groups = {}
            for spec in selections:
                if spec['column_name'] not in survey['columns']:
                    raise ValueError('Columna no presente en la encuesta')
                catalog = catalogs.get(spec['catalog_key'])
                if catalog is None or not catalog['compatible']:
                    raise ValueError('Catálogo inexistente o incompatible')
                groups.setdefault((spec['catalog_key'], spec['max_labels']), []).append(spec['column_name'])
            assignments = {}
            for (key, max_labels), columns in groups.items():
                catalog = catalogs[key]
                book = dict(id=identity(), project_id=project_id, name=catalog['field_id'],
                            categories=catalog['categories'], max_labels=max_labels, created_at=now())
                source = {name: catalog[name] for name in ('sheet', 'field_id', 'question_number', 'question_text', 'groupings')}
                source.update(filename=parsed['filename'], target_columns=columns)
                conn.execute(insert(db.codebooks).values(**book))
                conn.execute(insert(db.codebook_sources).values(codebook_id=book['id'], **source))
                record_activity(conn, actor_id, 'codebook_created', 'codebook', book['id'], project_id)
                record_activity(conn, actor_id, 'codebook_imported', 'codebook', book['id'], project_id)
                result['codebooks'].append({**book, 'source': source})
                for column in columns:
                    assignments[column] = dict(column_name=column, codebook_id=book['id'], question_text=catalog['question_text'])
            result['mappings'] = [assignments[spec['column_name']] for spec in selections]
        return result

    def create_book(self, project_id, data, actor_id=None):
        self.require(db.projects, project_id)
        item = dict(id=identity(), project_id=project_id, created_at=now(), **data.model_dump())
        with self.engine.begin() as conn:
            conn.execute(insert(db.codebooks).values(**item))
            record_activity(conn, actor_id, 'codebook_created', 'codebook', item['id'], project_id)
        return item

    def list_surveys(self, project_id):
        self.require(db.projects, project_id)
        with self.engine.connect() as conn:
            return [survey_public(dict(row)) for row in conn.execute(select(db.surveys).where(db.surveys.c.project_id == project_id).order_by(db.surveys.c.created_at.desc())).mappings()]

    def upload(self, project_id, filename, content, actor_id=None):
        self.require(db.projects, project_id)
        filename = PurePath(filename.replace('\\', '/')).name
        columns, rows = read_survey(filename, content)
        item = dict(id=identity(), project_id=project_id, filename=filename, original=content, rows=rows, columns=columns, total_records=len(rows), created_at=now())
        with self.engine.begin() as conn:
            conn.execute(insert(db.surveys).values(**item))
            record_activity(conn, actor_id, 'survey_uploaded', 'survey', item['id'], project_id)
        return survey_public(item)

    def create_job(self, data, provider_mode, actor_id=None):
        with self.engine.begin() as conn:
            survey = self.require(db.surveys, data.survey_id, conn)
            job_id, frozen_questions, planned, approved_records = identity(), [], [], []
            total = len(survey['rows']) * len(data.questions)
            for spec in data.questions:
                book = self.require(db.codebooks, spec.codebook_id, conn)
                if book['project_id'] != survey['project_id'] or spec.column_name not in survey['columns']:
                    raise ValueError('Columna o catálogo no pertenece a la encuesta/proyecto')
                question = dict(id=identity(), **spec.model_dump(), categories=book['categories'], max_labels=book['max_labels'])
                frozen_questions.append(question)
                code_col = find_code_column(spec.column_name, survey['columns'])
                texts = [row.get(spec.column_name, '') for row in survey['rows']]
                grouping = cluster_survey_responses(texts, similarity_threshold=85.0)
                groups = {}
                for index, row in enumerate(survey['rows']):
                    text = row.get(spec.column_name, '')
                    text = str(text) if text is not None else ''
                    precoded = extract_precoded_codes(row.get(code_col), book['categories'], book['max_labels']) if code_col else None
                    if precoded is not None:
                        approved_records.append(dict(id=identity(), job_id=job_id, question_id=question['id'], response_id=index, column_name=spec.column_name, text_original=text, codes=precoded, confidence=None, was_grouped=False, prototype_id=None, review_status='APPROVED', review_reason=None))
                    elif is_empty_or_symbol(text):
                        approved_records.append(dict(id=identity(), job_id=job_id, question_id=question['id'], response_id=index, column_name=spec.column_name, text_original=text, codes=['99'], confidence=None, was_grouped=False, prototype_id=None, review_status='APPROVED', review_reason=None))
                    else:
                        prototype = grouping.prototype_map[text.strip()]
                        groups.setdefault(prototype, []).append({'response_id':index, 'text':text})
                for text, members in groups.items():
                    planned.append(dict(id=identity(), ordinal=len(planned), job_id=job_id, question_id=question['id'], text=text, members=members, state='PENDING', result=None, attempts=0))
            timestamp = now()
            status = 'COMPLETED' if not planned else 'PENDING'
            job = dict(id=job_id, project_id=survey['project_id'], survey_id=survey['id'], name=data.name, status=status, questions=frozen_questions, total_records=total, processed_records=len(approved_records), provider_mode=provider_mode, pause_reason=None, pause_requested=False, metrics=initial_metrics(total, len(planned)), active_seconds=0, created_at=timestamp, updated_at=timestamp)
            conn.execute(insert(db.jobs).values(**job))
            for question in frozen_questions:
                conn.execute(insert(db.questions).values(id=question['id'], survey_id=survey['id'], column_name=question['column_name'], question_text=question['question_text'], codebook_id=question['codebook_id']))
            if planned:
                conn.execute(insert(db.prototypes), planned)
            if approved_records:
                conn.execute(insert(db.records), approved_records)
            if actor_id:
                conn.execute(insert(db.job_actors).values(job_id=job_id, actor_id=actor_id))
            record_activity(conn, actor_id, 'job_prepared', 'job', job_id, survey['project_id'], job_id)
        return self.get_job(job_id)

    def get_job(self, job_id):
        item = self.require(db.jobs, job_id)
        item['progress'] = round(item['processed_records'] / item['total_records'] * 100, 2) if item['total_records'] else 100
        return {key:value for key,value in item.items() if key not in {'pause_requested','active_seconds'}}

    def list_jobs(self, project_id=None):
        with self.engine.connect() as conn:
            query = select(db.jobs.c.id).order_by(db.jobs.c.created_at.desc())
            if project_id:
                query = query.where(db.jobs.c.project_id == project_id)
            ids = list(conn.scalars(query))
        return [self.get_job(id) for id in ids]

    def list_records(self, job_id, review_only=False, limit=200, offset=0):
        self.require(db.jobs, job_id)
        with self.engine.connect() as conn:
            query = select(db.records).where(db.records.c.job_id == job_id)
            if review_only:
                query = query.where(db.records.c.review_status == 'PENDING')
            return [dict(row) for row in conn.execute(query.order_by(db.records.c.response_id, db.records.c.column_name).offset(offset).limit(limit)).mappings()]

    def list_activity(self, job_id, limit=20):
        self.require(db.jobs, job_id)
        with self.engine.connect() as conn:
            query = (select(db.records)
                     .outerjoin(db.prototypes, db.records.c.prototype_id == db.prototypes.c.id)
                     .where(db.records.c.job_id == job_id)
                     .order_by(db.prototypes.c.ordinal.desc().nullslast(), db.records.c.response_id.desc())
                     .limit(limit))
            # The persisted plan is executed in ordinal order; vacíos are saved first.
            return list(reversed([dict(row) for row in conn.execute(query).mappings()]))
