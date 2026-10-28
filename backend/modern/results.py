"""Auditable human decisions, complete analytics and safe in-memory export."""
from collections import Counter
import csv
import io
from sqlalchemy import select, update, insert
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from . import db
from .repository import identity, now
from .engine import parse_codes


def lock_editable_job(repo, conn, id):
    changed = conn.execute(update(db.jobs).where(db.jobs.c.id == id, db.jobs.c.status != 'RUNNING').values(updated_at=now())).rowcount
    if not changed:
        repo.require(db.jobs, id, conn)
        raise HTTPException(409, 'Pause el trabajo antes de revisar')
    return repo.require(db.jobs, id, conn)


def approve(repo, job_id, record_id, codes):
    with repo.engine.begin() as conn:
        job = lock_editable_job(repo, conn, job_id)
        record = repo.require(db.records, record_id, conn)
        if record['job_id'] != job_id:
            raise KeyError('Recurso no encontrado')
        question = next(q for q in job['questions'] if q['id'] == record['question_id'])
        _, invalid = parse_codes(';'.join(codes), question['categories'], question['max_labels'])
        if invalid:
            raise ValueError('Seleccione códigos válidos del catálogo y respete el máximo permitido')
        conn.execute(update(db.records).where(db.records.c.id == record_id).values(codes=codes, review_status='APPROVED'))
        conn.execute(insert(db.audit).values(id=identity(), job_id=job_id, record_id=record_id, action='human_approved', payload={'before':record['codes'], 'after':codes}, created_at=now()))
        return repo.require(db.records, record_id, conn)


def batch(repo, job_id, data):
    ids = list(dict.fromkeys(data.record_ids))
    if len(ids) != len(data.record_ids):
        raise ValueError('Identificadores repetidos')
    with repo.engine.begin() as conn:
        job = lock_editable_job(repo, conn, job_id)
        items = [repo.require(db.records, id, conn) for id in ids]
        if any(item['job_id'] != job_id for item in items):
            raise KeyError('Recurso no encontrado')
        if len({item['question_id'] for item in items}) != 1:
            raise ValueError('Seleccione registros de una sola pregunta')
        frozen = [dict(q) for q in job['questions']]
        question = next(q for q in frozen if q['id'] == items[0]['question_id'])
        if data.code in {'77','88','99'} or data.code in {c['code'] for c in question['categories']}:
            raise ValueError('El nuevo código ya existe o está reservado')
        question['categories'] = question['categories'] + [{'code':data.code, 'label':data.label}]
        book = dict(conn.execute(select(db.codebooks).where(db.codebooks.c.id == question['codebook_id']).with_for_update()).mappings().one())
        if data.code in {c['code'] for c in book['categories']}:
            raise ValueError('El código ya existe en el catálogo actualizado')
        conn.execute(update(db.codebooks).where(db.codebooks.c.id == book['id']).values(categories=book['categories'] + [{'code':data.code,'label':data.label}]))
        conn.execute(update(db.jobs).where(db.jobs.c.id == job_id).values(questions=frozen))
        for item in items:
            conn.execute(update(db.records).where(db.records.c.id == item['id']).values(codes=[data.code], review_status='APPROVED'))
            conn.execute(insert(db.audit).values(id=identity(), job_id=job_id, record_id=item['id'], action='human_category_created', payload={'before':item['codes'],'after':[data.code],'label':data.label}, created_at=now()))
    return {'updated':len(ids)}


def analytics(repo, id):
    job = repo.get_job(id)
    counts, pending = Counter(), 0
    with repo.engine.connect() as conn:
        for row in conn.execute(select(db.records.c.question_id, db.records.c.codes, db.records.c.review_status).where(db.records.c.job_id == id)):
            counts.update((row.question_id, code) for code in row.codes)
            pending += row.review_status == 'PENDING'
    special = {'77':'Otro / nueva categoría','88':'Ninguno / satisfacción','99':'No sabe / no responde'}
    labels, columns = {}, {}
    for q in job['questions']:
        labels[q['id']] = special | {c['code']:c['label'] for c in q['categories']}
        columns[q['id']] = q['column_name']
    distribution = [{'question_id':question_id,'column_name':columns[question_id],'code':code,'label':labels[question_id].get(code,code),'count':count} for (question_id,code),count in counts.most_common()]
    return {'distribution':distribution, 'total_records':job['total_records'], 'processed_records':job['processed_records'], 'pending_review':pending}


def safe_cell(value):
    if isinstance(value, str) and (value.startswith(('\t','\r','\n')) or value.lstrip().startswith(('=','+','-','@'))):
        return "'" + value
    return value


def export(repo, id, format):
    job = repo.get_job(id)
    if job['status'] != 'COMPLETED':
        raise HTTPException(409, 'La exportación final requiere un trabajo completado')
    survey = repo.require(db.surveys, job['survey_id'])
    original_columns = survey['columns']
    added = [f"C{q['column_name']}_{i+1}" for q in job['questions'] for i in range(q['max_labels'])]
    if len(set(added)) != len(added) or set(added) & set(original_columns):
        raise HTTPException(409, 'Las columnas CNC entrarían en conflicto con columnas originales')

    orig_map = {}
    for idx, col in enumerate(original_columns):
        cleaned = str(col).strip().upper()
        if cleaned not in orig_map:
            orig_map[cleaned] = (idx + 1, col)

    existing_question_cols = {}
    for q in job['questions']:
        target = f"C{q['column_name']}".strip().upper()
        if target in orig_map:
            col_idx, orig_col_name = orig_map[target]
            existing_question_cols[q['id']] = (col_idx, orig_col_name)

    output = [dict(row) | {column:'' for column in added} for row in survey['rows']]
    questions = {q['id']:q for q in job['questions']}
    with repo.engine.connect() as conn:
        for record in conn.execute(select(db.records).where(db.records.c.job_id == id)).mappings():
            q = questions[record['question_id']]
            codes = record['codes']
            for i, code in enumerate(codes):
                output[record['response_id']][f"C{q['column_name']}_{i+1}"] = code
            if q['id'] in existing_question_cols:
                _, orig_name = existing_question_cols[q['id']]
                output[record['response_id']][orig_name] = ';'.join(codes) if codes else ''

    columns = original_columns + added
    if format == 'csv':
        stream = io.StringIO(newline='')
        writer = csv.writer(stream)
        writer.writerow([safe_cell(column) for column in columns])
        writer.writerows([safe_cell(row.get(key,'')) for key in columns] for row in output)
        return stream.getvalue().encode('utf-8-sig'), 'text/csv; charset=utf-8'
    if survey['filename'].lower().endswith(('.xlsx','.xls')):
        if survey['filename'].lower().endswith('.xlsx'):
            workbook = load_workbook(io.BytesIO(survey['original']))
        else:
            from .excel_xls import legacy_sheets
            workbook = Workbook()
            workbook.remove(workbook.active)
            for name, values in legacy_sheets(survey['original'], first_only=False):
                target = workbook.create_sheet(name)
                for values_row in values:
                    target.append([safe_cell(value) for value in values_row])
        for sheet in workbook:
            for row in sheet.iter_rows():
                for cell in row:
                    cell.value = safe_cell(cell.value)
        sheet = workbook.active
        for col_idx, orig_name in existing_question_cols.values():
            for row_index, row in enumerate(output, start=2):
                val = row.get(orig_name, '')
                if val:
                    cell = sheet.cell(row_index, col_idx, safe_cell(val))
                    cell.number_format = '@'
        for i, column in enumerate(added, start=len(original_columns)+1):
            sheet.cell(1,i,column)
        for row_index, row in enumerate(output,start=2):
            for i, column in enumerate(added,start=len(original_columns)+1):
                cell=sheet.cell(row_index,i,row[column])
                cell.number_format='@'
    else:
        workbook=Workbook()
        sheet=workbook.active
        sheet.title='Codificación CNC'
        sheet.append([safe_cell(column) for column in columns])
        for row in output:
            sheet.append([safe_cell(row.get(key,'')) for key in columns])
        for row in sheet.iter_rows(min_col=len(original_columns)+1,min_row=2):
            for cell in row:
                cell.number_format='@'
    stream=io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
