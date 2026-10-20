import io
import pytest
from conftest import SessionClient as TestClient
from openpyxl import Workbook
from test_study import make_study
from test_execution import FakeProvider, wait_job


def test_malformed_csv_returns_controlled_422(factory):
    with TestClient(factory()) as client:
        project=client.post('/api/v1/projects',json={'name':'Sintético','client':'QA'}).json()
        response=client.post(f"/api/v1/projects/{project['id']}/surveys",files={'file':('synthetic.csv',b'id,P3\n001,"unterminated','text/csv')})
        assert response.status_code==422
        assert client.get(f"/api/v1/projects/{project['id']}/surveys").json()=={'surveys':[]}


def test_excel_resources_close_when_iteration_fails(monkeypatch):
    from modern.study import read_survey
    import openpyxl
    workbook=Workbook(); workbook.active.append(['id','P3']); output=io.BytesIO();workbook.save(output);workbook.close()
    real_loader=openpyxl.load_workbook
    closed=[]
    def load(*args,**kwargs):
        book=real_loader(*args,**kwargs)
        close=book.close
        def closing():
            closed.append(True);close()
        book.close=closing
        def fail(*args,**kwargs):
            yield ('id','P3')
            raise SyntaxError('Malformed synthetic XML')
        book.active.iter_rows=fail
        return book
    monkeypatch.setattr(openpyxl,'load_workbook',load)
    with pytest.raises(ValueError,match='Excel inválido'):
        read_survey('synthetic.xlsx',output.getvalue())
    assert closed==[True]


def test_live_activity_is_latest_checkpoint_not_highest_row_number(factory):
    with TestClient(factory(provider=FakeProvider())) as client:
        *_,job=make_study(client,'id,P3\n001,crédito rápido\n002,nueva idea independiente\n003,crédito rápido\n004,crédito rápido\n')
        client.post(f"/api/v1/jobs/{job['id']}/start")
        assert wait_job(client,job['id'])['status']=='COMPLETED'
        response=client.get(f"/api/v1/jobs/{job['id']}/events")
        import json
        payload=json.loads(next(line[6:] for line in response.text.splitlines() if line.startswith('data: ')))
        assert payload['records'][-1]['text_original']=='nueva idea independiente'
