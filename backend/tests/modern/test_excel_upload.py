"""Only synthetic spreadsheets are created in memory; no external workbooks."""
import io
import datetime
import pytest
import xlwt
from modern.study import read_survey


def synthetic_xls():
    book=xlwt.Workbook()
    sheet=book.add_sheet('Respuestas')
    for column,value in enumerate(['id','P3','fecha','consentimiento']):sheet.write(0,column,value)
    for column,value in enumerate(['001','crédito rápido',datetime.datetime(2026,1,2),True]):
        sheet.write(1,column,value,xlwt.easyxf(num_format_str='YYYY-MM-DD') if column==2 else xlwt.Style.default_style)
    sheet.write(2,0,'002');sheet.write(2,1,'.')
    extra=book.add_sheet('Notas');extra.write(0,0,'Conservar hoja auxiliar');extra.write(1,0,'=texto no ejecutable')
    output=io.BytesIO();book.save(output);return output.getvalue()


def test_legacy_excel_upload_preserves_text_dates_and_booleans():
    columns,rows=read_survey('respuestas.XLS',synthetic_xls())
    assert columns==['id','P3','fecha','consentimiento']
    assert len(rows)==2
    assert rows[0]['id']=='001' and rows[0]['P3']=='crédito rápido'
    assert rows[0]['fecha']=='2026-01-02T00:00:00'
    assert rows[0]['consentimiento'] is True
    assert rows[1]['P3']=='.' and rows[1]['fecha']==''


def test_legacy_excel_export_preserves_all_sheet_values_and_cnc_codes(tmp_path):
    from modern.db import database
    from modern.study import StudyRepository
    from modern.schemas import CodebookCreate,JobCreate
    from modern import db,results
    from modern.repository import identity
    from sqlalchemy import insert,update
    from openpyxl import load_workbook
    engine=database(f'sqlite:///{(tmp_path / "xls-export.sqlite3").as_posix()}')
    repo=StudyRepository(engine)
    project=repo.create_project('Excel sintético','QA')
    survey=repo.upload(project['id'],'synthetic.xls',synthetic_xls())
    book=repo.create_book(project['id'],CodebookCreate(name='Catálogo',categories=[{'code':'01','label':'Crédito'}],max_labels=1))
    job=repo.create_job(JobCreate(survey_id=survey['id'],name='Export sintético',questions=[{'column_name':'P3','question_text':'Motivo','codebook_id':book['id']}]),'blocked')
    with engine.begin() as conn:
        conn.execute(update(db.jobs).where(db.jobs.c.id==job['id']).values(status='COMPLETED',processed_records=2))
        conn.execute(insert(db.records).values(id=identity(),job_id=job['id'],question_id=job['questions'][0]['id'],response_id=0,column_name='P3',text_original='crédito rápido',codes=['01'],confidence=None,was_grouped=False,review_status='APPROVED'))
    content,_=results.export(repo,job['id'],'xlsx')
    workbook=load_workbook(io.BytesIO(content))
    try:
        assert workbook.sheetnames==['Respuestas','Notas']
        assert workbook['Respuestas']['A2'].value=='001'
        assert workbook['Respuestas']['E1'].value=='CP3_1'
        assert workbook['Respuestas']['E2'].value=='01'
        assert workbook['Respuestas']['E3'].value=='99'
        assert workbook['Notas']['A1'].value=='Conservar hoja auxiliar'
        assert workbook['Notas']['A2'].value=="'=texto no ejecutable"
    finally:workbook.close();engine.dispose()


@pytest.mark.parametrize('content',[b'not an Excel workbook',b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'])
def test_corrupt_legacy_excel_is_controlled(content):
    with pytest.raises(ValueError,match='inválido'):
        read_survey('synthetic.xls',content)


def test_legacy_excel_resource_released_on_limit_failure(monkeypatch):
    from modern import excel_xls
    released=[]
    class FakeBook:
        nsheets=1
        def sheet_by_index(self,index):
            return type('Sheet',(),{'nrows':50002,'ncols':2})()
        def release_resources(self):released.append(True)
    monkeypatch.setattr(excel_xls.xlrd,'open_workbook',lambda **kwargs:FakeBook())
    with pytest.raises(ValueError,match='50.000'):
        read_survey('synthetic.xls',b'synthetic workbook bytes')
    assert released==[True]


def test_legacy_excel_bounds_cell_count_before_allocating_rows(monkeypatch):
    from modern import excel_xls
    monkeypatch.setattr(excel_xls,'MAX_CELLS',1)
    with pytest.raises(ValueError,match='celdas'):
        read_survey('synthetic.xls',synthetic_xls())
