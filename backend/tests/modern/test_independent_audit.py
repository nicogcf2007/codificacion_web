"""Independent v1 contract audit: synthetic data, fake clients, no remote I/O.

Only this new test module is written by the reviewer. Regression assertions are
not xfailed: a contract violation must keep the audit visibly red.
"""
import ast
import asyncio
from concurrent.futures import ThreadPoolExecutor
import csv
import io
from pathlib import Path
import sys
import threading
from types import SimpleNamespace

import pytest
from conftest import SessionClient as TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select, update

from modern import db, study
from modern.engine import Inference, OpenAIProvider, ProviderPaused, build_messages
from modern.schemas import BatchReview, Category, ReviewUpdate
from test_execution import FakeProvider, wait_job
from test_study import make_study


CATEGORIES = [{'code': '01', 'label': 'Créditos'}, {'code': '02', 'label': 'Atención'}]
SYNTHETIC_SECRET = 'synthetic-secret-never-real'


@pytest.fixture(autouse=True)
def audit_no_sdk_network(monkeypatch):
    # Prevent even accidental SDK construction from using a real transport.
    def forbidden_client(*args, **kwargs):
        raise AssertionError('A real OpenAI client is forbidden in this audit')
    monkeypatch.setitem(sys.modules, 'openai', SimpleNamespace(OpenAI=forbidden_client))
    for name in ('INPUT_USD_PER_MILLION', 'OUTPUT_USD_PER_MILLION', 'BASELINE_RECORD_USD'):
        monkeypatch.delenv(name, raising=False)


def fake_completion(raw='01'):
    return SimpleNamespace(
        choices=[SimpleNamespace(message=SimpleNamespace(content=raw))],
        usage=SimpleNamespace(prompt_tokens=10, completion_tokens=2),
    )


class FakeHTTPError(Exception):
    def __init__(self, status):
        self.status_code = status
        super().__init__('Synthetic provider body contains ' + SYNTHETIC_SECRET)


class FakeCompletions:
    def __init__(self, error=None, raw='01'):
        self.error = error
        self.raw = raw
        self.calls = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        if self.error is not None:
            raise self.error
        return fake_completion(self.raw)


def fake_openai(error=None, raw='01'):
    completions = FakeCompletions(error, raw)
    client = SimpleNamespace(chat=SimpleNamespace(completions=completions))
    return OpenAIProvider(client=client), completions


def set_headers(client):
    from conftest import login
    login(client)


def create_project(client, name='Audit'):
    result = client.post('/api/v1/projects', json={'name': name, 'client': 'Sintético'})
    assert result.status_code == 201, result.text
    return result.json()


def create_book(client, project, label='Créditos', max_labels=1):
    result = client.post(f"/api/v1/projects/{project['id']}/codebooks", json={
        'name': 'Audit catalog', 'categories': [{'code': '01', 'label': label}],
        'max_labels': max_labels,
    })
    assert result.status_code == 201, result.text
    return result.json()


def create_survey(client, project, data, filename='synthetic.csv'):
    result = client.post(f"/api/v1/projects/{project['id']}/surveys", files={
        'file': (filename, data, 'application/octet-stream'),
    })
    assert result.status_code == 201, result.text
    return result.json()


def create_job(client, survey, specs):
    result = client.post('/api/v1/coding/jobs', json={
        'survey_id': survey['id'], 'name': 'Independent audit', 'questions': specs,
    })
    assert result.status_code == 201, result.text
    return result.json()


def question(column, book, text='¿Por qué?'):
    return {'column_name': column, 'question_text': text, 'codebook_id': book['id']}


def finish_job(client, job):
    result = client.post(f"/api/v1/jobs/{job['id']}/start")
    assert result.status_code == 200, result.text
    result = wait_job(client, job['id'])
    assert result['status'] == 'COMPLETED', result
    return result


def prototype_rows(repo, job_id):
    with repo.engine.connect() as conn:
        return [dict(row) for row in conn.execute(
            select(db.prototypes).where(db.prototypes.c.job_id == job_id)
        ).mappings()]


def test_english_messages_identical_to_isolated_trusted_legacy_method():
    # AST isolation compiles ONLY the local pure method, never its module.
    # Therefore load_dotenv, global SDK clients and legacy caches cannot run.
    path = Path(__file__).resolve().parents[2] / 'core' / 'ai_engine.py'
    tree = ast.parse(path.read_text(encoding='utf-8'), filename=str(path))
    owner = next(node for node in tree.body if isinstance(node, ast.ClassDef)
                 and node.name == 'AICodingEngine')
    method = next(node for node in owner.body if isinstance(node, ast.FunctionDef)
                  and node.name == 'build_openai_messages')
    method.decorator_list = []
    method.returns = None
    for arg in method.args.args:
        arg.annotation = None
    module = ast.fix_missing_locations(ast.Module(body=[method], type_ignores=[]))
    namespace = {}
    exec(compile(module, str(path), 'exec'), namespace)
    original = namespace['build_openai_messages']
    for max_labels in (1, 2, 6):
        assert build_messages('¿Qué cambiaría?', 'rápido y amable', CATEGORIES, max_labels) == original(
            None, '¿Qué cambiaría?', 'rápido y amable', CATEGORIES,
            context='', max_labels=max_labels,
        )
    assert 'core.ai_engine' not in sys.modules


@pytest.mark.parametrize('status,uncertain', [(401, False), (429, False), (500, True), (503, True), (None, True)])
def test_fake_openai_errors_pause_without_codes_secrets_or_sdk_retry(factory, caplog, status, uncertain):
    error = FakeHTTPError(status) if status is not None else ConnectionError(SYNTHETIC_SECRET)
    provider, completions = fake_openai(error)
    with pytest.raises(ProviderPaused) as caught:
        provider.code('¿Por qué?', 'respuesta sintética', CATEGORIES, 2)
    assert caught.value.uncertain is uncertain
    assert SYNTHETIC_SECRET not in caught.value.reason
    assert len(completions.calls) == 1
    completions.calls.clear()
    app = factory(provider=provider)
    with TestClient(app) as client:
        set_headers(client)
        *_, job = make_study(client, 'id,P3\n001,crédito rápido\n')
        assert client.post(f"/api/v1/jobs/{job['id']}/start").status_code == 200
        paused = wait_job(client, job['id'])
        assert paused['status'] == 'PAUSED'
        assert paused['processed_records'] == 0
        assert paused['metrics']['api_calls'] == 1
        assert SYNTHETIC_SECRET not in str(paused)
        assert client.get(f"/api/v1/jobs/{job['id']}/records").json()['records'] == []
        prototypes = prototype_rows(app.state.repository, job['id'])
        assert [p['state'] for p in prototypes] == ['UNCERTAIN' if uncertain else 'PENDING']
        assert len(completions.calls) == 1
        if uncertain:
            assert client.post(f"/api/v1/jobs/{job['id']}/resume").status_code == 409
            assert len(completions.calls) == 1
        completions.error = None
        result = client.post(f"/api/v1/jobs/{job['id']}/resume", json={'retry_uncertain': uncertain})
        assert result.status_code == 200, result.text
        assert wait_job(client, job['id'])['status'] == 'COMPLETED'
        records = client.get(f"/api/v1/jobs/{job['id']}/records").json()['records']
        assert len(records) == 1 and records[0]['codes'] == ['01']
        with app.state.repository.engine.connect() as conn:
            actions = list(conn.scalars(select(db.audit.c.action).where(db.audit.c.job_id == job['id'])))
        assert actions == (['uncertain_retry_authorized'] if uncertain else [])
    assert SYNTHETIC_SECRET not in caplog.text


def test_removed_key_route_cannot_rotate_or_leak_provider_credentials(factory, monkeypatch, caplog):
    created = []
    def constructor(**kwargs):
        created.append(kwargs)
        return SimpleNamespace(chat=SimpleNamespace(completions=FakeCompletions()))
    monkeypatch.setitem(sys.modules, 'openai', SimpleNamespace(OpenAI=constructor))
    provider, old_client = fake_openai(FakeHTTPError(401))
    app = factory(provider=provider)
    with TestClient(app) as client:
        body = {'api_key': SYNTHETIC_SECRET}
        # Provider credentials can only be changed by server configuration.
        assert client.post('/api/v1/settings/openai-key', json=body).status_code == 404
        assert provider.client is not None
        *_, job = make_study(client, 'id,P3\n001,crédito rápido\n')
        assert client.post(f"/api/v1/jobs/{job['id']}/start").status_code == 200
        paused = wait_job(client, job['id'])
        assert paused['status'] == 'PAUSED' and paused['processed_records'] == 0
        invalid = client.post('/api/v1/settings/openai-key', json={'api_key': SYNTHETIC_SECRET * 30})
        assert invalid.status_code == 404
        assert SYNTHETIC_SECRET not in invalid.text
        assert SYNTHETIC_SECRET not in client.get('/api/v1/health').text
        assert provider.client is not None and provider.api_key is None
        old_client.error = None  # synthetic server-side provider recovery
        assert client.post(f"/api/v1/jobs/{job['id']}/resume").status_code == 200
        assert wait_job(client, job['id'])['status'] == 'COMPLETED'
        assert created == []
        assert len(old_client.calls) == 2
        repo = app.state.repository
        with repo.engine.connect() as conn:
            for table in (db.jobs, db.codebooks, db.audit, db.records, db.prototypes):
                assert SYNTHETIC_SECRET not in repr(list(conn.execute(select(table))))
    assert SYNTHETIC_SECRET not in caplog.text


@pytest.mark.parametrize('value', ['1', '001', ' 01', '01 ', '01\n', '01\r', '０１', '١٢', 1, True])
def test_two_digit_inputs_are_literal_never_coerced(value):
    from pydantic import ValidationError
    for schema, payload in (
        (Category, {'code': value, 'label': 'Audit'}),
        (ReviewUpdate, {'codes': [value]}),
        (BatchReview, {'record_ids': ['synthetic-id'], 'code': value, 'label': 'Audit'}),
    ):
        with pytest.raises(ValidationError):
            schema.model_validate(payload)
    assert Category(code='01', label='Audit').code == '01'


@pytest.mark.parametrize('value', ['true', 'yes', 1])
def test_uncertain_retry_requires_literal_boolean_consent(factory, value):
    app = factory(provider=FakeProvider())
    with TestClient(app) as client:
        *_, job = make_study(client, 'id,P3\n001,crédito rápido\n')
        with app.state.repository.engine.begin() as conn:
            conn.execute(update(db.jobs).where(db.jobs.c.id == job['id']).values(status='PAUSED'))
            conn.execute(update(db.prototypes).where(db.prototypes.c.job_id == job['id']).values(state='UNCERTAIN'))
        result = client.post(f"/api/v1/jobs/{job['id']}/resume", json={'retry_uncertain': value})
        assert result.status_code == 422, 'Non-boolean input must not authorize another possible charge'


def test_restart_preserves_checkpoint_plan_and_requires_consent(factory, monkeypatch):
    provider = FakeProvider()
    app = factory(provider=provider)
    with TestClient(app) as client:
        *_, job = make_study(client, 'id,P3\n001,crédito rápido\n002,categoría completamente nueva\n003,.\n')
        runtime = app.state.runtime
        first = runtime.next_prototype(job['id'])
        runtime.checkpoint(job['id'], first, Inference('01', 10, 2), 0.1)
        second = runtime.next_prototype(job['id'])
        with app.state.repository.engine.begin() as conn:
            conn.execute(update(db.jobs).where(db.jobs.c.id == job['id']).values(status='RUNNING'))
        before = prototype_rows(app.state.repository, job['id'])
    def no_reclustering(*args, **kwargs):
        raise AssertionError('Restart/resume must use the persisted plan')
    monkeypatch.setattr(study, 'cluster_survey_responses', no_reclustering)
    new_app = factory(provider=provider)
    with TestClient(new_app) as client:
        state = client.get(f"/api/v1/jobs/{job['id']}").json()
        assert state['status'] == 'PAUSED' and state['processed_records'] == 2
        assert client.post(f"/api/v1/jobs/{job['id']}/resume").status_code == 409
        assert provider.calls == []
        assert client.post(f"/api/v1/jobs/{job['id']}/resume", json={'retry_uncertain': True}).status_code == 200
        completed = wait_job(client, job['id'])
        assert completed['status'] == 'COMPLETED' and completed['processed_records'] == 3
        after = prototype_rows(new_app.state.repository, job['id'])
        assert {p['id'] for p in after} == {p['id'] for p in before}
        assert next(p for p in after if p['id'] == first['id'])['attempts'] == 1
        assert next(p for p in after if p['id'] == second['id'])['attempts'] == 2
        assert provider.calls == [second['text']]
        assert len(client.get(f"/api/v1/jobs/{job['id']}/records").json()['records']) == 3
        assert client.post(f"/api/v1/jobs/{job['id']}/resume", json={'retry_uncertain': True}).status_code == 409


def test_concurrent_starts_claim_once_and_running_review_is_rejected(factory):
    entered, release = threading.Event(), threading.Event()
    class Waiting(FakeProvider):
        def code(self, *args):
            entered.set()
            assert release.wait(8)
            return super().code(*args)
    provider = Waiting()
    app = factory(provider=provider)
    with TestClient(app) as client:
        *_, job = make_study(client, 'id,P3\n001,crédito rápido\n002,.\n')
        try:
            with ThreadPoolExecutor(max_workers=4) as pool:
                futures = [pool.submit(client.post, f"/api/v1/jobs/{job['id']}/start") for _ in range(4)]
                statuses = [future.result(timeout=5).status_code for future in futures]
            assert sorted(statuses) == [200, 409, 409, 409]
            assert entered.wait(3)
            existing = client.get(f"/api/v1/jobs/{job['id']}/records").json()['records'][0]
            assert client.patch(f"/api/v1/jobs/{job['id']}/records/{existing['id']}", json={'codes': ['01']}).status_code == 409
            assert client.post(f"/api/v1/jobs/{job['id']}/review/batch", json={
                'record_ids': [existing['id']], 'code': '03', 'label': 'Human',
            }).status_code == 409
        finally:
            release.set()
        assert wait_job(client, job['id'])['status'] == 'COMPLETED'
        assert provider.calls == ['crédito rápido']


def test_cancelled_inflight_is_uncertain_and_cannot_auto_repeat(factory):
    entered, release = threading.Event(), threading.Event()
    class Waiting(FakeProvider):
        def code(self, *args):
            entered.set()
            assert release.wait(8)
            return super().code(*args)
    app = factory(provider=Waiting())
    with TestClient(app) as client:
        *_, job = make_study(client, 'id,P3\n001,crédito rápido\n')
        try:
            assert client.post(f"/api/v1/jobs/{job['id']}/start").status_code == 200
            assert entered.wait(3)
            async def cancel_task():
                task = app.state.runtime.tasks[job['id']]
                task.cancel()
                await task
            client.portal.call(cancel_task)
            assert client.get(f"/api/v1/jobs/{job['id']}").json()['status'] == 'PAUSED'
            assert prototype_rows(app.state.repository, job['id'])[0]['state'] == 'UNCERTAIN'
            assert client.get(f"/api/v1/jobs/{job['id']}/records").json()['records'] == []
            assert client.post(f"/api/v1/jobs/{job['id']}/resume").status_code == 409
        finally:
            release.set()


def test_cross_study_creation_and_record_edits_are_rejected_atomically(factory):
    app = factory(provider=FakeProvider())
    with TestClient(app) as client:
        a, book_a, survey_a, job_a = make_study(client, 'id,P3\n001,crédito rápido\n')
        b, book_b, survey_b, job_b = make_study(client, 'id,P3\n001,otra respuesta\n')
        rejected = client.post('/api/v1/coding/jobs', json={
            'survey_id': survey_a['id'], 'name': 'Cross study',
            'questions': [question('P3', book_b)],
        })
        assert rejected.status_code == 422
        assert len(client.get('/api/v1/jobs').json()['jobs']) == 2
        finish_job(client, job_a)
        finish_job(client, job_b)
        record_a = client.get(f"/api/v1/jobs/{job_a['id']}/records").json()['records'][0]
        record_b = client.get(f"/api/v1/jobs/{job_b['id']}/records").json()['records'][0]
        assert client.patch(f"/api/v1/jobs/{job_a['id']}/records/{record_b['id']}", json={'codes': ['01']}).status_code == 404
        rejected = client.post(f"/api/v1/jobs/{job_a['id']}/review/batch", json={
            'record_ids': [record_a['id'], record_b['id']], 'code': '03', 'label': 'Cross study',
        })
        assert rejected.status_code == 404
        assert client.get(f"/api/v1/jobs/{job_a['id']}").json()['questions'][0]['categories'] == CATEGORIES
        with app.state.repository.engine.connect() as conn:
            assert conn.scalar(select(func.count()).select_from(db.audit)) == 0
        current = client.get(f"/api/v1/jobs/{job_a['id']}/records").json()['records'][0]
        assert current['codes'] == record_a['codes']


def test_multi_question_analytics_does_not_merge_different_labels_for_same_code(factory):
    app = factory(provider=FakeProvider())
    with TestClient(app) as client:
        project = create_project(client)
        first = create_book(client, project, 'Créditos')
        second = create_book(client, project, 'Amabilidad')
        survey = create_survey(client, project, b'id,P3,P4\n001,credito rapido,trato amable\n')
        job = create_job(client, survey, [question('P3', first), question('P4', second)])
        finish_job(client, job)
        distribution = client.get(f"/api/v1/jobs/{job['id']}/analytics").json()['distribution']
        assert {(row['code'], row['label'], row['count']) for row in distribution} == {
            ('01', 'Créditos', 1), ('01', 'Amabilidad', 1),
        }, 'Question-specific meanings must not be relabelled with the last codebook'


def test_batch_mixed_questions_is_atomic_and_snapshot_is_frozen(factory):
    app = factory(provider=FakeProvider())
    with TestClient(app) as client:
        project = create_project(client)
        first = create_book(client, project, 'Créditos')
        second = create_book(client, project, 'Amabilidad')
        survey = create_survey(client, project, b'id,P3,P4\n001,credito rapido,trato amable\n')
        job = create_job(client, survey, [question('P3', first), question('P4', second)])
        finish_job(client, job)
        records = client.get(f"/api/v1/jobs/{job['id']}/records").json()['records']
        response = client.post(f"/api/v1/jobs/{job['id']}/review/batch", json={
            'record_ids': [r['id'] for r in records], 'code': '03', 'label': 'Human',
        })
        assert response.status_code == 422
        assert client.get(f"/api/v1/jobs/{job['id']}/records").json()['records'] == records
        assert client.get(f"/api/v1/jobs/{job['id']}").json()['questions'] == job['questions']
        with app.state.repository.engine.begin() as conn:
            conn.execute(update(db.codebooks).where(db.codebooks.c.id == first['id']).values(categories=[{'code': '01', 'label': 'Changed elsewhere'}]))
        assert client.get(f"/api/v1/jobs/{job['id']}").json()['questions'] == job['questions']
        with app.state.repository.engine.connect() as conn:
            assert conn.scalar(select(func.count()).select_from(db.audit)) == 0


@pytest.mark.parametrize('format', ['csv', 'xlsx'])
def test_export_escapes_dangerous_original_column_headers(factory, format):
    with TestClient(factory(provider=FakeProvider())) as client:
        *_, job = make_study(client, 'id,P3,=1+1\n001,crédito rápido,normal\n')
        finish_job(client, job)
        response = client.get(f"/api/v1/jobs/{job['id']}/export?format={format}")
        assert response.status_code == 200
        if format == 'csv':
            header = next(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
            assert header[2].startswith("'"), 'CSV headers are spreadsheet cells too'
        else:
            workbook = load_workbook(io.BytesIO(response.content), data_only=False)
            try:
                formulas = [cell.coordinate for row in workbook.active for cell in row if cell.data_type == 'f']
                assert formulas == [], 'CSV-to-XLSX must not promote imported headers to formulas'
            finally:
                workbook.close()


@pytest.mark.parametrize('dangerous', ['=1+1', '+1+1', '-1+1', '@SUM(1)', '\t=1', '\r=1', '\n=1', '  =1+1'])
def test_csv_and_xlsx_export_escape_dangerous_values(factory, dangerous):
    source = io.StringIO(newline='')
    writer = csv.writer(source)
    writer.writerow(['id', 'P3', 'metadata'])
    writer.writerow(['001', 'crédito rápido', dangerous])
    with TestClient(factory(provider=FakeProvider())) as client:
        *_, job = make_study(client, source.getvalue())
        finish_job(client, job)
        response = client.get(f"/api/v1/jobs/{job['id']}/export?format=csv")
        rows = list(csv.DictReader(io.StringIO(response.content.decode('utf-8-sig'))))
        assert rows[0]['id'] == '001' and rows[0]['metadata'] == "'" + dangerous
        excel = client.get(f"/api/v1/jobs/{job['id']}/export?format=xlsx")
        workbook = load_workbook(io.BytesIO(excel.content))
        try:
            assert workbook.active['C2'].data_type == 's'
            # XML normalizes literal CR line endings when openpyxl reloads;
            # this assertion audits neutralization, not XML newline spelling.
            assert workbook.active['C2'].value in ("'" + dangerous, "'" + dangerous.replace('\r', '\n'))
            assert workbook.active['D2'].value == '01'
            assert workbook.active['D2'].number_format == '@'
        finally:
            workbook.close()


def test_xlsx_original_all_sheets_have_no_formulas_and_codes_keep_zeroes(factory):
    source = Workbook()
    sheet = source.active
    sheet.append(['id', 'P3', 'metadata'])
    sheet.append(['001', 'crédito rápido', '=1+1'])
    second = source.create_sheet('Other original sheet')
    second.append(['=1+1', 'unchanged'])
    stream = io.BytesIO()
    source.save(stream)
    source.close()
    with TestClient(factory(provider=FakeProvider())) as client:
        project = create_project(client)
        book = create_book(client, project)
        survey = create_survey(client, project, stream.getvalue(), 'synthetic.xlsx')
        job = create_job(client, survey, [question('P3', book)])
        finish_job(client, job)
        response = client.get(f"/api/v1/jobs/{job['id']}/export?format=xlsx")
        workbook = load_workbook(io.BytesIO(response.content), data_only=False)
        try:
            assert workbook.sheetnames == ['Sheet', 'Other original sheet']
            assert all(cell.data_type != 'f' for tab in workbook for row in tab for cell in row)
            assert workbook.active['A2'].value == '001'
            assert workbook.active['D2'].value == '01'
            assert workbook.active['D2'].number_format == '@'
            assert workbook['Other original sheet']['B1'].value == 'unchanged'
        finally:
            workbook.close()


def test_export_collision_refuses_without_overwriting_original(factory):
    app = factory(provider=FakeProvider())
    with TestClient(app) as client:
        *_, job = make_study(client, 'id,P3,CP3_1\n001,crédito rápido,original\n')
        finish_job(client, job)
        for format in ('csv', 'xlsx'):
            assert client.get(f"/api/v1/jobs/{job['id']}/export?format={format}").status_code == 409
        original = app.state.repository.require(db.surveys, job['survey_id'])
        assert original['rows'][0]['CP3_1'] == 'original'


def test_analytics_and_export_cover_all_records_beyond_ui_limit_and_multicodes(factory):
    class Multi(FakeProvider):
        def code(self, *args):
            self.calls.append(args[1])
            return Inference('01;02', 10, 2)
    app = factory(provider=Multi())
    data = 'id,P3\n' + ''.join(f'{index:04d},crédito rápido\n' for index in range(215))
    with TestClient(app) as client:
        *_, job = make_study(client, data)
        completed = finish_job(client, job)
        assert completed['total_records'] == completed['processed_records'] == 215
        assert len(client.get(f"/api/v1/jobs/{job['id']}/records").json()['records']) == 200
        stats = client.get(f"/api/v1/jobs/{job['id']}/analytics").json()
        assert {row['code']: row['count'] for row in stats['distribution']} == {'01': 215, '02': 215}
        assert stats['pending_review'] == 0
        response = client.get(f"/api/v1/jobs/{job['id']}/export?format=csv")
        rows = list(csv.DictReader(io.StringIO(response.content.decode('utf-8-sig'))))
        assert len(rows) == 215 and rows[-1]['id'] == '0214'
        assert all(row['CP3_1'] == '01' and row['CP3_2'] == '02' for row in rows)


def test_malformed_success_is_reviewed_not_silently_approved(factory):
    provider, calls = fake_openai(raw='secret-bearing explanation: ' + SYNTHETIC_SECRET)
    app = factory(provider=provider)
    with TestClient(app) as client:
        set_headers(client)
        *_, job = make_study(client, 'id,P3\n001,crédito rápido\n')
        finish_job(client, job)
        response = client.get(f"/api/v1/jobs/{job['id']}/records")
        record = response.json()['records'][0]
        assert record['codes'] == ['99'] and record['review_status'] == 'PENDING'
        assert 'inválida' in record['review_reason']
        assert SYNTHETIC_SECRET not in response.text
        assert record['confidence'] is None
        assert len(calls.calls) == 1


def test_checkpoint_database_error_rolls_back_records_and_requires_consent(factory):
    from sqlalchemy import event
    provider = FakeProvider()
    app = factory(provider=provider)
    with TestClient(app) as client:
        *_, job = make_study(client, 'id,P3\n001,crédito rápido\n')
        repo = app.state.repository
        failures = []
        def fail_final_checkpoint(conn, cursor, statement, parameters, context, executemany):
            if statement.startswith('UPDATE coding_jobs SET processed_records=') and not failures:
                failures.append(statement)
                raise RuntimeError('Synthetic checkpoint failure ' + SYNTHETIC_SECRET)
        event.listen(repo.engine, 'before_cursor_execute', fail_final_checkpoint)
        try:
            assert client.post(f"/api/v1/jobs/{job['id']}/start").status_code == 200
            paused = wait_job(client, job['id'])
            assert paused['status'] == 'PAUSED' and paused['processed_records'] == 0
            assert failures, 'Fault injection must hit after record inserts inside the transaction'
            assert SYNTHETIC_SECRET not in str(paused)
            assert client.get(f"/api/v1/jobs/{job['id']}/records").json()['records'] == []
            assert prototype_rows(repo, job['id'])[0]['state'] == 'UNCERTAIN'
            assert client.post(f"/api/v1/jobs/{job['id']}/resume").status_code == 409
        finally:
            event.remove(repo.engine, 'before_cursor_execute', fail_final_checkpoint)
        assert client.post(f"/api/v1/jobs/{job['id']}/resume", json={'retry_uncertain': True}).status_code == 200
        completed = wait_job(client, job['id'])
        assert completed['status'] == 'COMPLETED' and completed['processed_records'] == 1
        assert completed['metrics']['input_tokens'] == 10
        assert completed['metrics']['output_tokens'] == 2
        assert len(provider.calls) == 2
        assert len(client.get(f"/api/v1/jobs/{job['id']}/records").json()['records']) == 1


@pytest.mark.parametrize('kind', ['malformed_xlsx_xml', 'oversized_csv_field'])
def test_invalid_or_large_upload_is_a_controlled_client_error_not_500(factory, kind):
    if kind == 'malformed_xlsx_xml':
        from zipfile import ZipFile, ZIP_DEFLATED
        workbook = Workbook()
        workbook.active.append(['id', 'P3'])
        workbook.active.append(['001', 'crédito rápido'])
        source = io.BytesIO()
        workbook.save(source)
        workbook.close()
        corrupt = io.BytesIO()
        with ZipFile(io.BytesIO(source.getvalue())) as archive, ZipFile(corrupt, 'w', ZIP_DEFLATED) as output:
            for name in archive.namelist():
                data = b'<worksheet invalid' if name == 'xl/worksheets/sheet1.xml' else archive.read(name)
                output.writestr(name, data)
        data, filename = corrupt.getvalue(), 'synthetic.xlsx'
    else:
        data = ('id,P3,metadata\n001,crédito rápido,' + 'x' * 150000 + '\n').encode()
        filename = 'synthetic.csv'
    app = factory(provider=FakeProvider())
    with TestClient(app, raise_server_exceptions=False) as client:
        project = create_project(client)
        response = client.post(f"/api/v1/projects/{project['id']}/surveys", files={'file': (filename, data)})
        if kind == 'malformed_xlsx_xml':
            assert response.status_code == 422, 'Malformed office XML must have a controlled validation error'
        else:
            assert response.status_code in (201, 422), 'An undocumented csv field limit must not become HTTP 500'
        if response.status_code != 201:
            assert client.get(f"/api/v1/projects/{project['id']}/surveys").json()['surveys'] == []
