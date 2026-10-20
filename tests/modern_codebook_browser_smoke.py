"""End-to-end Excel workflow using only synthetic records and an isolated server.

Uses the actual cookie/CSRF API, a new scratch SQLite database and DemoProvider.
Does not load any workbook from excel, operational DB or provider credential.
"""
from __future__ import annotations
import io
import json
import os
from pathlib import Path
import secrets
import socket
import sys
import tempfile
import threading
import time
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))
from modern.app import create_app
from modern.engine import DemoProvider
from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright, expect
import uvicorn


def base_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Respuestas'
    sheet.append(['id', 'CP3_1', 'P3_1', 'CP3_2', 'P3_2', 'P33T'])
    sheet.append(['0001', '', 'crédito rápido', '', 'atención cordial', 'atención'])
    sheet.append(['0002', '', '.', '', 'crédito rápido', '-'])
    book.create_sheet('Notas').append(['Hoja auxiliar sintética'])
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def codebook_bytes():
    book = Workbook()
    book.active.title = 'Identificación'
    book.active.append(['Documento sintético; esta hoja no es un catálogo'])
    sheet = book.create_sheet('Codificación')
    sheet.append(['Id campo', 'Cod', 'Label', 'Agrupación', '# Pregunta del formulario', 'Nombre de la Pregunta'])
    sheet.append(['CP3_1-CP3_2', '01', 'Crédito', 'Servicio', '3', '¿Por qué recomendaría el servicio?'])
    sheet.append(['CP3_1-CP3_2', '02', 'Atención', 'Servicio', '3', None])
    sheet.append(['CP3_1-CP3_2', '99', 'No sabe / no responde', None, '3', None])
    sheet.append(['CP33T', '01', 'Atención', None, '33', '¿Qué se puede mejorar?'])
    sheet.append(['CP33T', '99', 'No sabe / no responde', None, '33', None])
    sheet.append(['CT32T_1', '001', 'Categoría incompatible', None, '32', 'Código de tres dígitos'])
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def run():
    scratch = Path(os.environ.get('TMPDIR') or (Path.home() / 'AppData/Local/hermes/cache/scratch')).resolve()
    output = Path(tempfile.mkdtemp(prefix='cnc-codebook-browser-', dir=scratch))
    # No inherited production credentials, DSN or auth setup can enter this test.
    for key in ('DATABASE_URL', 'OPENAI_API_KEY', 'CODING_PROVIDER', 'CNC_ADMIN_EMAIL',
                'CNC_ADMIN_PASSWORD', 'CNC_ADMIN_NAME', 'CNC_ENV', 'APP_ENV',
                'CNC_COOKIE_SECURE', 'CORS_ORIGINS'):
        os.environ.pop(key, None)
    password = secrets.token_urlsafe(24)
    app = create_app('sqlite:///' + (output / 'qa.sqlite3').as_posix(), DemoProvider(),
                     admin_email='admin@example.test', admin_password=password,
                     admin_name='Admin Sintético')
    listener = socket.socket()
    listener.bind(('127.0.0.1', 0))
    listener.listen(128)
    url = 'http://127.0.0.1:' + str(listener.getsockname()[1])
    server = uvicorn.Server(uvicorn.Config(app, log_level='error', access_log=False))
    worker = threading.Thread(target=server.run, kwargs={'sockets': [listener]}, daemon=True)
    worker.start()
    deadline = time.monotonic() + 15
    while not server.started:
        if time.monotonic() > deadline or not worker.is_alive():
            raise RuntimeError('Isolated QA server did not start')
        time.sleep(.05)
    errors, external, checks = [], [], []
    try:
        with sync_playwright() as toolkit:
            browser = toolkit.chromium.launch(headless=True)
            context = browser.new_context(viewport={'width':1440,'height':1000}, reduced_motion='reduce', accept_downloads=True)
            def guard(route):
                if urlparse(route.request.url).hostname not in {'127.0.0.1', 'localhost'}:
                    external.append('External request blocked')
                    route.abort()
                else:
                    route.continue_()
            context.route('**/*', guard)
            page = context.new_page()
            page.on('pageerror', lambda error: errors.append(str(error)))
            page.goto(url, wait_until='networkidle')
            page.get_by_label('Correo electrónico', exact=True).fill('admin@example.test')
            page.get_by_label('Contraseña', exact=True).fill(password)
            page.get_by_role('button', name='Entrar', exact=True).click()
            expect(page.get_by_role('heading', name='Espacio de proyectos')).to_be_visible()
            page.get_by_label('Nombre del estudio', exact=True).fill('Excel sintético')
            page.get_by_label('Cliente', exact=True).fill('QA aislada')
            with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/api/v1/projects')) as created:
                page.get_by_role('button', name='Crear proyecto', exact=True).click()
            assert created.value.status == 201
            project_id = created.value.json()['id']
            expect(page.get_by_role('heading', name='Preparar codificación')).to_be_visible()
            with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/surveys')) as uploaded:
                page.get_by_label('Subir base de respuestas', exact=True).set_input_files({'name':'base-sintetica.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':base_bytes()})
            assert uploaded.value.status == 201
            survey_id = uploaded.value.json()['id']
            for column in ('P3_1', 'P3_2', 'P33T'):
                page.locator('.column-option').filter(has=page.get_by_text(column, exact=True)).locator('input').check()
            expect(page.get_by_role('button', name='Preparar ejecución', exact=True)).to_be_disabled()
            with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/codebooks/preview')) as preview:
                page.get_by_label('Subir libro de códigos', exact=True).set_input_files({'name':'codigos-sinteticos.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':codebook_bytes()})
            assert preview.value.status == 200
            parsed = preview.value.json()
            assert len(parsed['catalogs']) == 3
            assert sum(item['compatible'] for item in parsed['catalogs']) == 2
            invalid = next(item for item in parsed['catalogs'] if not item['compatible'])
            assert invalid['categories'][0]['code'] == '001'
            assert context.request.get(url + f'/api/v1/projects/{project_id}/codebooks').json()['codebooks'] == [], 'Preview persisted a catalogue'
            expect(page.get_by_role('button', name='Preparar ejecución', exact=True)).to_be_disabled()
            page.screenshot(path=str(output / '01-correspondencias-oscuro.png'), full_page=True)
            with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/codebooks/import')) as imported:
                page.get_by_role('button', name='Confirmar correspondencias', exact=True).click()
            assert imported.value.status == 201
            loaded = imported.value.json()
            assert len(loaded['mappings']) == 3
            assert len(loaded['codebooks']) == 2, 'Shared catalogue duplicated'
            assert set(mapping['column_name'] for mapping in loaded['mappings']) == {'P3_1','P3_2','P33T'}
            assert loaded['mappings'][0]['question_text']
            books = context.request.get(url + f'/api/v1/projects/{project_id}/codebooks').json()['codebooks']
            assert len(books) == 2
            assert all(book['source']['filename']=='codigos-sinteticos.xlsx' for book in books)
            assert all(book['source']['sheet']=='Codificación' for book in books)
            assert set(books[0]['categories'][0]['code']) <= set('0123456789')
            expect(page.get_by_role('button', name='Preparar ejecución', exact=True)).to_be_enabled()
            page.screenshot(path=str(output / '02-confirmado-oscuro.png'), full_page=True)
            checks.extend(['separate base and codebook Excel uploads', 'three explicitly selected columns', 'auxiliary sheet ignored', 'one catalogue proposed for two columns', 'P3 does not match P33', '001 preserved and incompatible group excluded', 'preview is read-only', 'confirmation persists official categories and question text'])
            # Changing the selected set revokes the previous confirmation.
            target = page.locator('.column-option').filter(has=page.get_by_text('P33T', exact=True)).locator('input')
            target.uncheck()
            expect(page.get_by_role('button', name='Preparar ejecución', exact=True)).to_be_disabled()
            target.check()
            expect(page.get_by_role('button', name='Preparar ejecución', exact=True)).to_be_disabled()
            page.get_by_role('button', name='Confirmar correspondencias', exact=True).click()
            expect(page.get_by_role('button', name='Preparar ejecución', exact=True)).to_be_enabled()
            checks.append('selection changes invalidate confirmation')
            page.set_viewport_size({'width':390,'height':844})
            page.get_by_role('button', name='Tema claro', exact=True).click()
            assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth'), 'Mobile overflow'
            page.screenshot(path=str(output / '03-correspondencias-movil-claro.png'), full_page=True)
            page.set_viewport_size({'width':1440,'height':1000})
            with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/coding/jobs')) as prepared:
                page.get_by_role('button', name='Preparar ejecución', exact=True).click()
            assert prepared.value.status == 201
            job = prepared.value.json()
            assert job['survey_id'] == survey_id
            assert len(job['questions']) == 3 and job['total_records'] == 6
            assert len(context.request.get(url + f"/api/v1/jobs/{job['id']}/records").json()['records']) == 2
            page.get_by_role('button', name='Iniciar codificación', exact=True).click()
            expect(page.get_by_text('Completado', exact=True)).to_be_visible(timeout=30000)
            completed = context.request.get(url + f"/api/v1/jobs/{job['id']}").json()
            assert completed['processed_records'] == 6
            page.locator('.app-sidebar').get_by_role('button', name='Resultados', exact=True).click()
            with page.expect_download() as downloaded:
                page.get_by_role('button', name='Excel', exact=True).click()
            destination = output / 'export-sintetico.xlsx'
            downloaded.value.save_as(str(destination))
            exported = load_workbook(destination, data_only=False)
            assert exported.sheetnames == ['Respuestas','Notas']
            sheet = exported['Respuestas']
            assert sheet['A2'].value == '0001'
            headers = [cell.value for cell in sheet[1]]
            assert {'P3_1','P3_2','P33T'} <= set(headers)
            code_headers = [header for header in headers if str(header).startswith('C')]
            assert code_headers
            assert any(cell.value=='99' for row in sheet.iter_rows(min_row=2) for cell in row), 'Empty response code missing'
            exported.close()
            events = context.request.get(url + '/api/v1/admin/audit?limit=200').json()
            imports = [event for event in events['events'] if event['action']=='codebook_imported']
            assert imports and all(event['actor_email']=='admin@example.test' for event in imports)
            assert not errors, errors
            assert not external, external
            checks.extend(['390px light layout without overflow', 'multi-question preparation and synthetic demo completion', 'empty responses 99', 'Excel download preserves auxiliary sheet and leading zero ID', 'import audit readback with authenticated actor', 'no pageerrors or external requests'])
            result = dict(passed=True, checks=checks, provider='explicit synthetic demo only', records=6,
                          screenshots=[str(path) for path in sorted(output.glob('*.png'))], output=str(output))
            (output / 'browser-results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
            print(json.dumps(result,ensure_ascii=False,indent=2))
            browser.close()
    finally:
        server.should_exit = True
        worker.join(timeout=10)
        listener.close()


if __name__ == '__main__':
    run()
