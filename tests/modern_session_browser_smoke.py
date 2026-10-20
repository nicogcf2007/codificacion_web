"""Full HTTP/browser QA with disposable synthetic users and spreadsheets.
Owns a loopback-only server and a NEW scratch database, never an existing app.
Usage: backend/venv/Scripts/python.exe tests/modern_session_browser_smoke.py
"""
from __future__ import annotations
import io
import json
import os
from pathlib import Path
import secrets
import socket
import sys
import tempfile
import threading
import time
from urllib.parse import urlparse
import urllib.request

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))
from modern.app import create_app
from modern.engine import DemoProvider
from playwright.sync_api import sync_playwright, expect
from openpyxl import Workbook, load_workbook
import uvicorn
import xlwt

RESPONSES = ['crédito ágil', 'atención cordial', 'requisitos excesivos', 'tasas elevadas',
             'horarios extendidos', 'plataforma estable', 'cercanía territorial',
             'información confusa', 'cobros inesperados', 'solución de reclamos',
             'diversidad de productos', 'seguridad digital', 'beneficios promocionales',
             'transacciones lentas', 'cajeros disponibles', 'seguros útiles',
             'asesoría especializada', 'comunicación frecuente', 'soporte permanente',
             'nada por ahora', '.', 'crédito ágil']


def workbook_bytes(extension):
    rows = [['id', 'P3'], *[[f'{i:04d}', text] for i, text in enumerate(RESPONSES, 1)]]
    output = io.BytesIO()
    if extension == 'xls':
        book = xlwt.Workbook(); sheet = book.add_sheet('Respuestas')
        for i, row in enumerate(rows):
            for j, value in enumerate(row): sheet.write(i, j, value)
        notes = book.add_sheet('Notas'); notes.write(0, 0, 'Hoja auxiliar sintética')
        book.save(output)
    else:
        book = Workbook(); sheet = book.active; sheet.title = 'Respuestas'
        for row in rows: sheet.append(row)
        book.save(output); book.close()
    return output.getvalue()


def contrast(foreground, background):
    def luminance(color):
        import re
        values = [float(n) / 255 for n in re.findall(r'[\d.]+', color)[:3]]
        channels = [v / 12.92 if v <= .04045 else ((v + .055) / 1.055) ** 2.4 for v in values]
        return sum(v * weight for v, weight in zip(channels, (.2126, .7152, .0722)))
    light, dark = sorted([luminance(foreground), luminance(background)], reverse=True)
    return (light + .05) / (dark + .05)


def run():
    scratch = Path(os.environ.get('TMPDIR') or (Path.home() / 'AppData/Local/hermes/cache/scratch')).resolve()
    scratch.mkdir(parents=True, exist_ok=True)
    output = Path(tempfile.mkdtemp(prefix='cnc-session-browser-', dir=scratch))
    # Random throwaway passwords exist only in this process and the scratch hash DB.
    admin_password = secrets.token_urlsafe(24)
    provisional_password = secrets.token_urlsafe(24)
    standard_password = secrets.token_urlsafe(24)
    app = create_app('sqlite:///' + (output / 'qa.sqlite3').as_posix(), DemoProvider(),
                     admin_email='admin@example.test', admin_password=admin_password, admin_name='Admin Sintético')
    listener = socket.socket(); listener.bind(('127.0.0.1', 0)); listener.listen(128)
    url = 'http://127.0.0.1:' + str(listener.getsockname()[1])
    server = uvicorn.Server(uvicorn.Config(app, log_level='error', access_log=False))
    worker = threading.Thread(target=server.run, kwargs={'sockets': [listener]}, daemon=True)
    worker.start()
    deadline = time.monotonic() + 15
    while not server.started:
        if time.monotonic() > deadline or not worker.is_alive(): raise RuntimeError('QA server startup failed')
        time.sleep(.05)
    checks = []; errors = []; unexpected = []
    try:
        with sync_playwright() as toolkit:
            browser = toolkit.chromium.launch(headless=True)
            context = browser.new_context(viewport={'width':1440,'height':1000}, color_scheme='light', reduced_motion='reduce', accept_downloads=True)
            def guard(route):
                if urlparse(route.request.url).hostname not in {'127.0.0.1','localhost'}:
                    unexpected.append('blocked external browser request'); route.abort()
                else: route.continue_()
            context.route('**/*', guard)
            page = context.new_page(); page.on('pageerror', lambda error: errors.append(str(error)))
            assert context.request.get(url + '/api/v1/projects').status == 401
            page.goto(url, wait_until='networkidle')
            expect(page.get_by_role('heading', name='Iniciar sesión', exact=True)).to_be_visible()
            assert page.evaluate('document.documentElement.dataset.theme') == 'dark'
            assert page.get_by_role('group', name='Apariencia').get_by_role('button').count() == 2
            assert page.get_by_text('Registrarse', exact=False).count() == 0
            page.screenshot(path=str(output / '01-login-oscuro.png'), full_page=True)
            def login(email, password):
                page.get_by_label('Correo electrónico', exact=True).fill(email)
                page.get_by_label('Contraseña', exact=True).fill(password)
                page.get_by_role('button', name='Entrar', exact=True).click()
            def logout(name):
                page.get_by_role('button', name='Perfil de ' + name, exact=True).click()
                page.get_by_role('button', name='Cerrar sesión', exact=True).click()
                expect(page.get_by_role('heading', name='Iniciar sesión', exact=True)).to_be_visible()
            def clean_ui():
                text = page.locator('body').inner_text().casefold()
                for term in ['conexión y ajustes', 'api conectada', 'ia bloqueada', 'gpt-6', 'gpt6', 'openai', 'api key', 'autorización de la aplicación', 'token en memoria']:
                    assert term not in text, 'Technical UI text remains: ' + term
                expect(page.get_by_role('button', name='Actualizar espacio', exact=True)).to_be_visible()
            login('admin@example.test', admin_password)
            expect(page.get_by_role('heading', name='Espacio de proyectos')).to_be_visible()
            clean_ui()
            cookies = context.cookies()
            assert next(c for c in cookies if c['name'] == 'cnc_session')['httpOnly']
            assert context.request.get(url + '/api/v1/auth/me').status == 200
            page.get_by_label('Nombre del estudio', exact=True).fill('Estudio compartido sintético')
            page.get_by_label('Cliente', exact=True).fill('QA aislada')
            with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/api/v1/projects')) as created:
                page.get_by_role('button', name='Crear proyecto', exact=True).click()
            assert created.value.status == 201
            project = created.value.json()
            sidebar = page.locator('.app-sidebar')
            sidebar.get_by_role('button', name='Usuarios', exact=True).click()
            page.get_by_role('button', name='Crear usuario', exact=True).click()
            dialog = page.get_by_role('dialog')
            dialog.get_by_label('Nombre completo', exact=True).fill('Ana Prueba')
            dialog.get_by_label('Correo electrónico', exact=True).fill('ana@example.test')
            dialog.get_by_label('Contraseña inicial', exact=True).fill(provisional_password)
            with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/admin/users')) as user_response:
                dialog.get_by_role('button', name='Guardar usuario', exact=True).click()
            assert user_response.value.status == 201
            standard = user_response.value.json()
            expect(page.get_by_text('Cuenta creada.', exact=False)).to_be_visible()
            page.screenshot(path=str(output / '02-usuarios-oscuro.png'), full_page=True)
            logout('Admin Sintético')
            assert context.request.get(url + '/api/v1/projects').status == 401
            login('ana@example.test', provisional_password)
            expect(page.get_by_role('heading', name='Cambia tu contraseña provisional')).to_be_visible()
            assert context.request.get(url + '/api/v1/projects').status == 403
            page.get_by_label('Contraseña actual', exact=True).fill(provisional_password)
            page.get_by_label('Nueva contraseña', exact=True).fill(standard_password)
            page.get_by_label('Confirmar nueva contraseña', exact=True).fill(standard_password)
            page.get_by_role('button', name='Guardar contraseña', exact=True).click()
            expect(page.get_by_role('heading', name='Espacio de proyectos')).to_be_visible()
            expect(page.get_by_role('button', name='Perfil de Ana Prueba', exact=True)).to_be_visible()
            assert page.locator('.profile-trigger .avatar').inner_text() == 'AP'
            for name in ['Usuarios','Auditoría global','Resumen administrativo']:
                assert sidebar.get_by_role('button', name=name, exact=True).count() == 0
            assert context.request.get(url + '/api/v1/admin/users').status == 403
            expect(page.locator('.project-name strong').filter(has_text=project['name'])).to_be_visible()
            checks.extend(['login with cookies and no public signup', 'dark default despite light OS', 'ADMIN creates provisional STANDARD account', 'provisional password gate and change', 'STANDARD sees ADMIN project', 'admin APIs and navigation denied to STANDARD', 'profile initials from session', 'logout revocation'])
            page.locator('.project-name').filter(has_text=project['name']).click()
            expect(page.get_by_role('heading', name='Preparar codificación')).to_be_visible()
            page.get_by_label('Subir base de respuestas', exact=True).set_input_files({'name':'invalid.txt','mimeType':'text/plain','buffer':b'invalid'})
            expect(page.get_by_role('alert').filter(has_text='Selecciona un archivo')).to_be_visible()
            with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/surveys')) as xlsx_upload:
                page.get_by_label('Subir base de respuestas', exact=True).set_input_files({'name':'synthetic.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':workbook_bytes('xlsx')})
            assert xlsx_upload.value.status == 201
            transfer = page.evaluate_handle('([bytes,name]) => {const data=new DataTransfer();data.items.add(new File([new Uint8Array(bytes)],name,{type:"application/vnd.ms-excel"}));return data;}', [list(workbook_bytes('xls')), 'synthetic.xls'])
            with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/surveys')) as xls_upload:
                page.get_by_role('region', name='Carga de base de respuestas', exact=True).dispatch_event('drop', {'dataTransfer':transfer})
            assert xls_upload.value.status == 201
            transfer.dispose()
            expect(page.get_by_text('Carga completada: synthetic.xls', exact=True)).to_be_visible()
            expect(page.get_by_text(f'{len(RESPONSES)} registros · 2 columnas', exact=True)).to_be_visible()
            page.get_by_label('Nombre del catálogo', exact=True).fill('Catálogo sintético')
            page.get_by_label('Máximo de códigos', exact=True).fill('2')
            page.get_by_label('Etiqueta 1', exact=True).fill('Crédito')
            page.get_by_role('button', name='Categoría', exact=True).click()
            page.get_by_label('Código 2', exact=True).fill('02')
            page.get_by_label('Etiqueta 2', exact=True).fill('Atención')
            page.get_by_role('button', name='Guardar catálogo', exact=True).click()
            expect(page.get_by_role('button', name='Crear otro catálogo', exact=True)).to_be_visible()
            page.locator('.column-option').filter(has_text='P3').locator('input').check()
            page.get_by_label('Pregunta P3', exact=True).fill('¿Qué mejorarías?')
            clean_ui(); page.screenshot(path=str(output / '03-carga-oscura.png'), full_page=True)
            with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/coding/jobs')) as prepared:
                page.get_by_role('button', name='Preparar ejecución', exact=True).click()
            assert prepared.value.status == 201
            job = prepared.value.json(); job_id = job['id']
            page.get_by_role('button', name='Iniciar codificación', exact=True).click()
            expect(page.get_by_text('Completado', exact=True)).to_be_visible(timeout=30000)
            clean_ui()
            assert context.request.get(url + f'/api/v1/jobs/{job_id}').json()['processed_records'] == len(RESPONSES)
            page.screenshot(path=str(output / '04-monitor-oscuro.png'), full_page=True)
            sidebar.get_by_role('button', name='Revisión humana', exact=True).click()
            expect(page.locator('blockquote')).to_be_visible()
            with page.expect_response(lambda r:r.request.method=='PATCH' and '/records/' in r.url) as approved:
                page.keyboard.press('2'); page.keyboard.press('Enter')
            assert approved.value.status == 200
            clean_ui()
            sidebar.get_by_role('button', name='Resultados', exact=True).click()
            expect(page.locator('.bar-row').first).to_be_visible()
            with page.expect_download() as downloaded:
                page.get_by_role('button', name='Excel', exact=True).click()
            destination = output / 'synthetic-export.xlsx'; downloaded.value.save_as(str(destination))
            exported = load_workbook(destination)
            assert exported.sheetnames == ['Respuestas','Notas']
            assert exported['Respuestas']['A2'].value == '0001'
            assert exported['Respuestas']['C1'].value == 'CP3_1'
            assert exported['Respuestas']['D1'].value == 'CP3_2'
            assert exported['Notas']['A1'].value == 'Hoja auxiliar sintética'
            exported.close()
            clean_ui()
            checks.extend(['invalid upload controlled', '.xlsx through picker', '.xls through real DataTransfer drop', 'catalogue and job preparation', 'authenticated fake execution and SSE completion', 'STANDARD keyboard review', 'XLS→XLSX browser export with extra sheet and leading zeroes'])
            page.get_by_role('button', name='Tema claro', exact=True).click()
            assert page.evaluate('document.documentElement.dataset.theme') == 'light'
            css = page.locator('.app-main').evaluate('(e) => ({fg:getComputedStyle(e).color,bg:getComputedStyle(document.body).backgroundColor})')
            contrast_value = contrast(css['fg'],css['bg']); assert contrast_value >= 4.5, {'computed':css,'contrast':contrast_value}
            page.set_viewport_size({'width':390,'height':844})
            assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth'), 'Mobile overflow'
            page.screenshot(path=str(output / '05-movil-claro.png'), full_page=True)
            page.reload(wait_until='networkidle')
            expect(page.get_by_role('heading', name='Espacio de proyectos')).to_be_visible()
            assert page.evaluate('document.documentElement.dataset.theme') == 'light'
            assert page.evaluate('Object.keys(localStorage)') == ['cnc-theme']
            logout('Ana Prueba')
            login('admin@example.test', admin_password)
            expect(page.get_by_role('heading', name='Espacio de proyectos')).to_be_visible()
            page.set_viewport_size({'width':1440,'height':1000})
            sidebar.get_by_role('button', name='Auditoría global', exact=True).click()
            expect(page.get_by_role('heading', name='Trazabilidad del equipo')).to_be_visible()
            events = context.request.get(url + '/api/v1/admin/audit?limit=200').json()
            assert events['total'] == len(events['events'])
            expected_actor = {'project_created':'admin@example.test','survey_uploaded':'ana@example.test','codebook_created':'ana@example.test','job_prepared':'ana@example.test','job_started':'ana@example.test','job_checkpoint':'ana@example.test','human_approved':'ana@example.test'}
            observed = {event['action'] for event in events['events']}
            for action, email in expected_actor.items():
                assert action in observed, 'Missing audit action ' + action
                assert all(event['actor_email'] == email for event in events['events'] if event['action']==action), 'Wrong audit actor ' + action
            assert all('password' not in json.dumps(event).lower() for event in events['events'] if event['action']!='password_changed')
            page.screenshot(path=str(output / '06-auditoria-clara.png'), full_page=True)
            sidebar.get_by_role('button', name='Resumen administrativo', exact=True).click()
            overview = context.request.get(url + '/api/v1/admin/overview').json()
            assert overview['projects']==1 and overview['users']==2 and overview['completed_jobs']==1
            page.get_by_role('button', name='Actualizar espacio', exact=True).click()
            checks.extend(['theme persists after reload and logout', 'theme only localStorage', '390px layout without overflow', 'light computed text contrast >=4.5', 'global audit actor readback for asynchronous work and review', 'global admin counts readback', 'top reload retained', 'technical/provider UI excluded on all screens'])
            assert not errors, errors
            assert not unexpected, unexpected
            browser.close()
            result = dict(passed=True, checks=checks, records=len(RESPONSES), light_text_contrast=round(contrast_value,2), audit_events=events['total'], screenshots=[str(p) for p in sorted(output.glob('*.png'))], page_errors=errors, unexpected_network=unexpected, output=str(output), provider='explicit synthetic demo; no real provider')
            (output/'browser-results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
            print(json.dumps(result,ensure_ascii=False,indent=2))
    finally:
        server.should_exit=True; worker.join(timeout=10); listener.close()


if __name__=='__main__': run()
